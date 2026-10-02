import html, importlib.util, io, json, re, threading, time, traceback
from pathlib import Path
import streamlit as st

st.set_page_config(page_title="AutoMinutes", page_icon="🎬", layout="wide", initial_sidebar_state="expanded")

# --- GLASSMORPHISM TOKENS ---
T = dict(
    bg="#090D16",
    surface="rgba(255, 255, 255, 0.04)",
    field="rgba(255, 255, 255, 0.03)",
    line="rgba(255, 255, 255, 0.12)",
    ink="#F8FAFC",
    muted="#94A3B8",
    a1="#38BDF8",
    a2="#818CF8",
    ok="#34D399",
    danger="#F87171"
)
ROOT = ":root, .stApp{" + "".join(f"--{k}:{v};" for k, v in T.items()) + "}"

CSS = """
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@600;700;800&family=Inter:wght@400;500;600&display=swap');

/* Typography Split: Inter for Body, Plus Jakarta Sans for Headers */
html, body, .stApp,
.stApp *:not([data-testid="stIconMaterial"]):not(.material-icons):not(.material-symbols-rounded):not([data-testid="collapsedControl"] *):not([data-testid="stSidebarCollapseButton"] *) {
    font-family: 'Inter', sans-serif !important;
}

/* Fix raw ligature text leak on the sidebar toggle */
[data-testid="stIconMaterial"], .material-icons, .material-symbols-rounded,
[data-testid="collapsedControl"] span, [data-testid="stSidebarCollapseButton"] span {
    font-family: 'Material Symbols Rounded', 'Material Icons' !important;
    font-size: 1.25rem !important;
}

h1, h2, h3, h4, .hero, .logo, .kpi b, .chat-head, .pstage, .ring-in b, .step h4, .result-title {
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    letter-spacing: -0.02em;
}

.stApp, .stApp p, .stApp li, .stApp label, .stApp span, .stApp h1, .stApp h2, .stApp h3, .stApp h4 {
    color: var(--ink);
}
.stApp { background: var(--bg) !important; }

/* Frosted Ambient Background directly on container (no pseudo-elements intercepting pointer events) */
[data-testid="stAppViewContainer"] {
    background: radial-gradient(1000px 600px at 10% -10%, rgba(56, 189, 248, 0.15), transparent),
                radial-gradient(900px 500px at 90% 10%, rgba(129, 140, 248, 0.12), transparent),
                radial-gradient(800px 600px at 50% 100%, rgba(56, 189, 248, 0.08), transparent) !important;
    background-attachment: fixed !important;
}

#MainMenu, footer { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; pointer-events: none; }
header[data-testid="stHeader"] [data-testid="stToolbar"] { pointer-events: auto; }

.block-container { position: relative; z-index: 1; max-width: 1180px; padding-top: 2.4rem; padding-bottom: 4rem; }
.nav { display: flex; align-items: center; justify-content: space-between; margin-bottom: 2.2rem; }
.logo { display: flex; align-items: center; gap: .7rem; font-weight: 700; font-size: 1.15rem; }
.mark {
    width: 36px; height: 36px; border-radius: 10px; display: flex; align-items: center; justify-content: center; gap: 3px;
    background: transparent; border: 1px solid var(--line);
}
.mark i { width: 3px; border-radius: 2px; background: var(--ink); }

.pill {
    display: inline-flex; align-items: center; gap: .5rem; padding: .38rem .85rem;
    border: 1px solid var(--line); border-radius: 999px; background: rgba(255, 255, 255, 0.03);
    font-size: 0.82rem; font-weight: 600; color: var(--muted) !important;
}
.pill i { width: 8px; height: 8px; border-radius: 50%; background: var(--c); }
.pill.run i { animation: ping 1.4s infinite; }
@keyframes ping { 0%{box-shadow:0 0 0 0 color-mix(in srgb,var(--c) 70%,transparent)} 80%,100%{box-shadow:0 0 0 9px transparent} }

.hero { font-weight: 800; font-size: 2.7rem; line-height: 1.15; margin: 0 0 .6rem; }
.hero span {
    background: linear-gradient(90deg, #38BDF8, #818CF8);
    -webkit-background-clip: text; background-clip: text; -webkit-text-fill-color: transparent;
}
.sub { color: var(--muted) !important; font-size: 1rem; line-height: 1.6; max-width: 620px; margin: 0 0 1.8rem; }

/* Panels without isolated composite stacking bugs */
[data-testid="stVerticalBlockBorderWrapper"], [data-testid="stExpander"] details {
    background: var(--surface) !important;
    border: 1px solid var(--line) !important;
    border-radius: 16px !important;
    box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.36);
}

.stTextInput input, .stSelectbox div[data-baseweb="select"] > div, [data-testid="stFileUploaderDropzone"] {
    background: var(--field) !important;
    color: var(--ink) !important;
    border: 1px solid var(--line) !important;
    border-radius: 12px !important;
}
.stTextInput input { padding: .75rem .9rem; }
.stTextInput input::placeholder { color: var(--muted) !important; }
.stTextInput input:focus { border-color: var(--a1) !important; box-shadow: 0 0 0 2px rgba(56, 189, 248, 0.25) !important; }

[data-testid="stFileUploaderDropzone"] { border-style: dashed !important; }
[data-testid="stFileUploaderDropzone"]:hover { border-color: var(--a1) !important; }

/* Transparent Glass Buttons */
:root { --r: 12px; --h: 44px; }
.stButton > button, .stDownloadButton > button, [data-testid="stSidebar"] .stButton > button {
    width: 100%; height: var(--h); min-height: var(--h); padding: 0 1.2rem !important;
    border: 1px solid var(--line) !important;
    border-radius: var(--r) !important;
    font-weight: 600; font-size: 0.92rem;
    justify-content: center !important; text-align: center; white-space: nowrap;
    color: var(--ink) !important;
    background: transparent !important;
    box-shadow: none !important;
    transition: all .2s ease;
}
.stButton > button p, .stDownloadButton > button p {
    color: var(--ink) !important; margin: 0; font-weight: 600 !important; white-space: nowrap;
}
.stButton > button:hover, .stDownloadButton > button:hover {
    border-color: var(--a1) !important;
    color: var(--a1) !important;
    background: rgba(255, 255, 255, 0.05) !important;
    box-shadow: 0 0 16px rgba(56, 189, 248, 0.2) !important;
    transform: translateY(-1px);
}
.stButton > button:hover p, .stDownloadButton > button:hover p {
    color: var(--a1) !important;
}
.stButton > button:active, .stDownloadButton > button:active { transform: scale(.98); }
.stButton > button:disabled { opacity: .4; transform: none; }

.stButton > button[kind="primary"] {
    border: 1px solid var(--a1) !important;
    background: transparent !important;
    box-shadow: 0 0 14px rgba(56, 189, 248, 0.15) !important;
}
.stButton > button[kind="primary"] p { color: var(--a1) !important; }
.stButton > button[kind="primary"]:hover {
    background: rgba(56, 189, 248, 0.08) !important;
    box-shadow: 0 0 20px rgba(56, 189, 248, 0.3) !important;
}

.st-key-stop button {
    border-color: var(--danger) !important;
    color: var(--danger) !important;
    background: transparent !important;
}
.st-key-stop button p { color: var(--danger) !important; }
.st-key-stop button:hover {
    background: rgba(248, 113, 113, 0.08) !important;
    box-shadow: 0 0 18px rgba(248, 113, 113, 0.25) !important;
}

[data-testid="stSidebar"] {
    background: rgba(13, 18, 27, 0.95) !important;
    border-right: 1px solid var(--line);
}
[data-testid="stSidebar"] .stButton > button { justify-content: flex-start !important; text-align: left !important; }

[class*="st-key-del_"] button {
    width: var(--h) !important; min-width: var(--h) !important; padding: 0 !important;
    border: 1px solid var(--line) !important; background: transparent !important;
}
[class*="st-key-del_"] button:hover {
    color: var(--danger) !important; border-color: var(--danger) !important;
    background: rgba(248, 113, 113, 0.08) !important;
}

/* === FIT ALL 5 TABS & ABSOLUTELY PREVENT SCROLL CLICK FREEZING === */
.stTabs { 
    position: relative !important; 
    z-index: 1000 !important; 
}
.stTabs [data-baseweb="tab-list"] {
    display: flex !important;
    width: 100% !important;
    gap: 4px !important;
    background: rgba(9, 13, 22, 0.95) !important;
    border-bottom: 1px solid var(--line) !important;
    padding: 6px !important;
    border-radius: 10px;
    overflow-x: auto !important;
    white-space: nowrap !important;
    position: sticky !important;
    top: 0 !important;
    z-index: 1001 !important;
    pointer-events: auto !important;
}
.stTabs [data-baseweb="tab"] {
    flex: 1 1 0px !important;
    min-width: 90px !important;
    text-align: center !important;
    padding: 8px 10px !important;
    font-size: 0.88rem !important;
    font-weight: 600 !important;
    background: transparent !important;
    border-radius: 8px !important;
    border: none !important;
    cursor: pointer !important;
    pointer-events: auto !important;
}
.stTabs [data-baseweb="tab"] p {
    color: var(--muted) !important;
    margin: 0 !important;
    font-size: 0.88rem !important;
    pointer-events: none !important;
}
.stTabs [aria-selected="true"] {
    background: rgba(255, 255, 255, 0.08) !important;
}
.stTabs [aria-selected="true"] p {
    color: var(--ink) !important;
    font-weight: 700 !important;
}
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] { display: none !important; }

/* Safe Scroll Area with strict pointer-events isolation */
.stTabs [data-baseweb="tab-panel"] {
    padding-top: 0.8rem !important;
    position: relative !important;
    z-index: 100 !important;
    pointer-events: auto !important;
}

.transcript-scroll-area {
    max-height: 420px;
    overflow-y: scroll !important;
    overflow-x: hidden !important;
    padding: 1rem 1.2rem;
    background: rgba(0, 0, 0, 0.35);
    border: 1px solid var(--line);
    border-radius: 12px;
    font-size: 0.94rem;
    line-height: 1.75;
    white-space: pre-wrap;
    pointer-events: auto !important;
    overscroll-behavior: contain !important;
}

/* Custom Scrollbars */
::-webkit-scrollbar { width: 6px; height: 6px; }
::-webkit-scrollbar-track { background: rgba(0, 0, 0, 0.2); border-radius: 8px; }
::-webkit-scrollbar-thumb { background: rgba(255, 255, 255, 0.25); border-radius: 8px; }
::-webkit-scrollbar-thumb:hover { background: rgba(56, 189, 248, 0.5); }

.tab-action-row { display: flex; justify-content: flex-end; margin-bottom: 0.5rem; }
.glass-copy-btn {
    width: 32px; height: 32px; display: inline-flex; align-items: center; justify-content: center;
    border: 1px solid var(--line); border-radius: 8px; background: transparent; color: var(--muted);
    cursor: pointer; transition: all .2s ease;
}
.glass-copy-btn:hover { color: var(--a1); border-color: var(--a1); background: rgba(56, 189, 248, 0.08); }

.kpis { display: grid; grid-template-columns: repeat(3, 1fr); gap: .8rem; margin-bottom: 1rem; }
.kpi { border: 1px solid var(--line); background: var(--surface); border-radius: 14px; padding: .8rem 1rem; }
.kpi small { display: block; color: var(--muted); font-size: 0.72rem; font-weight: 700; text-transform: uppercase; }
.kpi b { font-size: 1.25rem; }

.chat-head { font-weight: 700; font-size: 1.2rem; }
.chat-sub { color: var(--muted) !important; font-size: .9rem; margin: .1rem 0 .8rem; }
[data-testid="stChatMessage"] { background: transparent; }
[data-testid="stChatInput"], [data-testid="stChatInput"] > div {
    background: var(--field) !important; border-radius: 14px; border-color: var(--line) !important;
}

.pcard { display: flex; flex-wrap: wrap; align-items: center; gap: 1.8rem; padding: 1.4rem .4rem .4rem; }
.ring {
    --p: 0; width: 120px; height: 120px; border-radius: 50%; flex: none; display: grid; place-items: center;
    background: conic-gradient(var(--a1), var(--a2) calc(var(--p)*1%), var(--line) calc(var(--p)*1%));
    box-shadow: 0 0 24px rgba(56, 189, 248, 0.2);
}
.ring-in { width: 96px; height: 96px; border-radius: 50%; background: var(--bg); display: flex; align-items: baseline; justify-content: center; padding-top: 30px; box-sizing: border-box; }
.ring-in b { font-weight: 800; font-size: 1.9rem; }
.ring-in small { color: var(--muted) !important; font-weight: 700; margin-left: 2px; }
.pinfo { flex: 1; min-width: 260px; }
.pstage { font-weight: 600; font-size: 1.15rem; }
.psub { color: var(--muted) !important; font-size: .88rem; margin: .15rem 0 .8rem; }
.pbar { height: 6px; border-radius: 99px; background: var(--line); overflow: hidden; }
.pfill { height: 100%; border-radius: 99px; transition: width .6s ease; background: linear-gradient(90deg, var(--a1), var(--a2)); }

.stages { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap: .55rem .9rem; margin-top: 1.1rem; }
.stg { display: flex; align-items: center; gap: .55rem; font-size: .88rem; font-weight: 600; color: var(--muted) !important; }
.stg .mk { width: 18px; height: 18px; border-radius: 50%; border: 1px solid var(--line); flex: none; display: inline-flex; align-items: center; justify-content: center; font-size: .65rem; }
.stg.done { color: var(--ink) !important; } .stg.done .mk { background: var(--ok); border-color: var(--ok); color: #000; }
.stg.active { color: var(--ink) !important; } .stg.active .mk { border-color: var(--a1); border-top-color: transparent; animation: spin .8s linear infinite; }
@keyframes spin { to { transform: rotate(360deg); } }

.steps { display: grid; grid-template-columns: repeat(auto-fit, minmax(230px, 1fr)); gap: 1rem; margin-top: 1.6rem; }
.step {
    border: 1px solid var(--line); background: var(--surface); border-radius: 16px;
    padding: 1.3rem 1.4rem; transition: all .25s ease;
}
.step:hover { transform: translateY(-4px); border-color: var(--a1); box-shadow: 0 12px 30px rgba(56, 189, 248, 0.12); }
.step b {
    display: inline-flex; width: 30px; height: 30px; align-items: center; justify-content: center;
    border-radius: 8px; border: 1px solid var(--a1); color: var(--a1); margin-bottom: .8rem;
}
.step h4 { margin: 0 0 .3rem; font-size: 1.05rem; }
.step p { margin: 0; color: var(--muted) !important; font-size: .92rem; line-height: 1.55; }
"""
st.markdown(f"<style>{ROOT}{CSS}</style>", unsafe_allow_html=True)

st.session_state.setdefault("result", None)
st.session_state.setdefault("messages", [])
st.session_state.setdefault("took", 0)
st.session_state.setdefault("aid", "x")

@st.cache_resource
def _preload():
    def run():
        try:
            import main  # noqa: F401
        except Exception:
            pass
    threading.Thread(target=run, daemon=True).start()
    return True

_preload()

STAGES = [("Prepare audio", 0), ("Transcribe", 15), ("Summarise", 62), ("Action items", 76),
          ("Key decisions", 84), ("Open questions", 91), ("Chat index", 96)]

class Job:
    def __init__(self):
        self.progress, self.stage = 0.0, "Starting"
        self.status = "running"
        self.cancel = threading.Event()
        self.result, self.error, self.collected = None, "", False
        self.detail, self.source = "", {}
        self.started, self.took = time.time(), 0.0

def _worker(job: Job, source: str, language: str):
    try:
        job.progress, job.stage = 1, "Loading AI models"
        from main import run_pipeline
        def on_progress(pct, msg):
            job.progress, job.stage = pct, msg
        job.result = run_pipeline(source, language, on_progress=on_progress, should_stop=job.cancel.is_set)
        if not job.cancel.is_set():
            job.took, job.status = time.time() - job.started, "done"
    except Exception as e:
        if not job.cancel.is_set():
            job.status, job.error, job.detail = "error", friendly_error(e), traceback.format_exc()

DATA = Path(__file__).parent / "data"
HIST, UPLOADS = DATA / "history", DATA / "uploads"
for _d in (HIST, UPLOADS):
    _d.mkdir(parents=True, exist_ok=True)
HAS_DOCX = importlib.util.find_spec("docx") is not None
HAS_PDF = importlib.util.find_spec("reportlab") is not None
_BUL = re.compile(r"^(?:[-*•]|\d+[.)])\s+")
STOP = set("the a an and or of to in on for with that this is are was were be it as at by from about what who when how "
           "why which will would can could should did does have has had not you your they their we our".split())

def clean_title(t) -> str:
    return re.sub(r"[*`#]+", "", str(t)).strip().strip("\"'“” ").strip()

def valid_youtube(u: str) -> bool:
    return bool(re.match(r"^https?://(www\.|m\.|music\.)?(youtube\.com|youtu\.be)/\S+", u.strip()))

def friendly_error(e) -> str:
    m = str(e).lower()
    if any(k in m for k in ("private video", "unavailable", "sign in", "403", "yt_dlp", "downloaderror", "unable to download")):
        return "Couldn't download this video. It may be private, age-restricted or blocked. Try another link or upload the file."
    if any(k in m for k in ("api key", "api_key", "authentication", "401")):
        return "There's a problem with your API key. Check the keys in your .env file."
    if any(k in m for k in ("rate limit", "429", "quota")):
        return "The AI service is rate-limiting requests. Wait a minute and try again, or check your plan's quota."
    if "ffmpeg" in m or "ffprobe" in m:
        return "ffmpeg wasn't found. Install it and make sure it is on your PATH."
    return "Something went wrong while processing. See the technical details below."

def save_upload(f) -> str:
    path = UPLOADS / f"{int(time.time())}_{re.sub(r'[^A-Za-z0-9._-]', '_', f.name)}"
    path.write_bytes(f.getbuffer())
    return str(path)

def save_record(r: dict, took: float) -> str:
    aid = time.strftime("%Y%m%d-%H%M%S")
    rec = {k: v for k, v in r.items() if k != "rag_chain"}
    rec.update(id=aid, took=took, created=time.strftime("%d %b, %H:%M"))
    (HIST / f"{aid}.json").write_text(json.dumps(rec), encoding="utf-8")
    return aid

@st.cache_data(show_spinner=False)
def _titles(stamp: tuple) -> list:
    out = []
    for name, _ in stamp:
        try:
            d = json.loads((HIST / name).read_text(encoding="utf-8"))
            out.append({"id": d["id"], "title": clean_title(d.get("title", "Untitled")), "created": d.get("created", "")})
        except Exception:
            pass
    return out

def list_records() -> list:
    return _titles(tuple((p.name, p.stat().st_mtime_ns) for p in sorted(HIST.glob("*.json"), reverse=True)))

def open_record(aid: str):
    d = json.loads((HIST / f"{aid}.json").read_text(encoding="utf-8"))
    d["title"] = clean_title(d.get("title", "Untitled"))
    d["rag_chain"] = None
    st.session_state.update(result=d, aid=aid, took=d.get("took", 0), messages=[], job=None)

def delete_record(aid: str):
    p = HIST / f"{aid}.json"
    try:
        src = json.loads(p.read_text(encoding="utf-8")).get("source") or {}
        if src.get("type") == "file":
            Path(src["value"]).unlink(missing_ok=True)
    except Exception:
        pass
    p.unlink(missing_ok=True)

def prune_uploads():
    used = set()
    for p in HIST.glob("*.json"):
        try:
            used.add((json.loads(p.read_text(encoding="utf-8")).get("source") or {}).get("value"))
        except Exception:
            pass
    for p in UPLOADS.glob("*"):
        if str(p) not in used and time.time() - p.stat().st_mtime > 86400:
            p.unlink(missing_ok=True)

if "_pruned" not in st.session_state:
    prune_uploads()
    st.session_state["_pruned"] = True

def parse_items(text: str) -> list:
    lines = [l.strip() for l in text.splitlines() if l.strip() and not l.strip().startswith("#")]
    bullets = [_BUL.sub("", l) for l in lines if _BUL.match(l)]
    return bullets or [l for l in lines if not l.endswith(":")]

def clean_for_copy(text: str) -> str:
    if not text:
        return ""
    cleaned = re.sub(r'#{1,6}\s*', '', text)
    cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', cleaned)
    cleaned = re.sub(r'\*(.*?)\*', r'\1', cleaned)
    cleaned = re.sub(r'`(.*?)`', r'\1', cleaned)
    return cleaned

def render_copy_icon(text: str, key_id: str):
    clean_txt = clean_for_copy(text)
    safe_payload = html.escape(json.dumps(clean_txt))
    icon_svg = '<svg width="15" height="15" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2" ry="2"></rect><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"></path></svg>'
    check_svg = '<svg width="15" height="15" viewBox="0 0 24 24" fill="none" stroke="#34D399" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><polyline points="20 6 9 17 4 12"></polyline></svg>'
    
    html_btn = (
        f'<div class="tab-action-row">'
        f'<button class="glass-copy-btn" id="btn_{key_id}" title="Copy" data-text="{safe_payload}" '
        f'onclick=\'(function(b){{navigator.clipboard.writeText(JSON.parse(b.getAttribute("data-text")));'
        f'b.innerHTML="{check_svg}";setTimeout(function(){{b.innerHTML="{icon_svg}";}}, 1500);}})(this);\'>'
        f'{icon_svg}'
        f'</button>'
        f'</div>'
    )
    st.markdown(html_btn, unsafe_allow_html=True)

def typewriter(text: str):
    for tok in re.findall(r"\S+\s*", text):
        yield tok
        time.sleep(0.015)

def _tok(t: str) -> set:
    return {w for w in re.findall(r"[a-z0-9']+", t.lower()) if len(w) > 2 and w not in STOP}

def find_sources(query: str, segs: list, k: int = 2) -> list:
    q, scored = _tok(query), []
    for s in segs:
        t = _tok(s["text"])
        if t:
            scored.append((len(q & t) / (len(t) ** 0.5), s))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [{"t": s.get("start"), "text": s["text"][:280]} for sc, s in scored[:k] if sc > 0]

def passages(text: str, n: int = 70) -> list:
    w = text.split()
    return [{"start": None, "text": " ".join(w[i:i + n])} for i in range(0, len(w), n)]

def show_sources(srcs):
    if not srcs:
        return
    with st.expander(f"Sources ({len(srcs)})"):
        for s in srcs:
            st.caption(s["text"])

def sections_of(r: dict) -> tuple:
    return (("Summary", r["summary"]), ("Action items", r["action_items"]),
            ("Key decisions", r["key_decisions"]), ("Open questions", r["open_questions"]))

def _clean(line: str) -> str:
    return re.sub(r"[*_`]+", "", line).lstrip("# ").strip()

@st.cache_data(show_spinner=False)
def build_docx(title: str, secs: tuple, transcript: str) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading(title, 0)
    for name, text in secs:
        d.add_heading(name, 1)
        for line in filter(None, (l.strip() for l in text.splitlines())):
            d.add_paragraph(_clean(_BUL.sub("", line)), style="List Bullet" if _BUL.match(line) else None)
    d.add_heading("Transcript", 1)
    d.add_paragraph(transcript)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

@st.cache_data(show_spinner=False)
def build_pdf(title: str, secs: tuple, transcript: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    ss, buf = getSampleStyleSheet(), io.BytesIO()
    fix = lambda t: html.escape(t.encode("cp1252", "replace").decode("cp1252"))
    story = [Paragraph(fix(title), ss["Title"])]
    for name, text in secs:
        story.append(Paragraph(fix(name), ss["Heading2"]))
        for line in filter(None, (l.strip() for l in text.splitlines())):
            story.append(Paragraph(("• " if _BUL.match(line) else "") + fix(_clean(_BUL.sub("", line))), ss["BodyText"]))
    story += [Paragraph("Transcript", ss["Heading2"]), Paragraph(fix(transcript), ss["BodyText"])]
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()

def nav_html(state: str) -> str:
    label, color, cls = {
        "idle": ("Ready", "#94A3B8", ""), "run": ("Analysing", "#38BDF8", "run"), "done": ("Complete", "#34D399", ""),
        "stopped": ("Stopped", "#FBBF24", ""), "error": ("Error", "#F87171", ""),
    }[state]
    bars = "".join(f'<i style="height:{h}px"></i>' for h in (8, 16, 22, 14, 9))
    return (f'<div class="nav"><div class="logo"><div class="mark">{bars}</div>AutoMinutes</div>'
            f'<div class="pill {cls}" style="--c:{color}"><i></i>{label}</div></div>')

def progress_html(pct: float, stage: str, elapsed: float) -> str:
    p = max(0, min(100, int(pct)))
    active = max(i for i, (_, s) in enumerate(STAGES) if pct >= s)
    rows = ""
    for i, (name, _) in enumerate(STAGES):
        state = "done" if (i < active or p >= 100) else "active" if i == active else "todo"
        rows += f'<div class="stg {state}"><span class="mk">{"✓" if state == "done" else ""}</span><span>{name}</span></div>'
    mm, ss = divmod(int(elapsed), 60)
    return (
        f'<div class="pcard"><div class="ring" style="--p:{p}"><div class="ring-in"><b>{p}</b><small>%</small></div></div>'
        f'<div class="pinfo"><div class="pstage">{html.escape(stage)}</div><div class="psub">Elapsed {mm:02d}:{ss:02d}</div>'
        f'<div class="pbar"><div class="pfill" style="width:{p}%"></div></div><div class="stages">{rows}</div></div></div>'
    )

with st.sidebar:
    st.markdown('<div class="chat-head">History</div>', unsafe_allow_html=True)
    recs = list_records()
    if not recs:
        st.caption("Your Analyses will appear here.")
    for rec in recs:
        c1, c2 = st.columns([4, 1])
        c1.button(f"{rec['title'][:30]} · {rec['created']}", key=f"open_{rec['id']}",
                  on_click=open_record, args=(rec["id"],), use_container_width=True)
        c2.button("×", key=f"del_{rec['id']}", help="Delete", on_click=delete_record, args=(rec["id"],), use_container_width=True)

nav_slot = st.empty()
st.markdown(
    '<h1 class="hero">Turn recordings into <span>decisions</span>.</h1>'
    '<p class="sub">Paste a YouTube link or upload a recording. Get a transcript, summary, action items and '
    "open questions, then chat with the content.</p>",
    unsafe_allow_html=True,
)

job = st.session_state.get("job")
running = bool(job and job.status == "running")

def begin():
    ss = st.session_state
    up = ss.get("up")
    chosen = ss.get("_up_path") if up else (ss.get("url") or "").strip()
    if not chosen:
        ss.warn = "Paste a YouTube link or upload a file first."
    elif not up and not valid_youtube(chosen):
        ss.warn = "That doesn't look like a YouTube link. Paste a full youtube.com or youtu.be URL, or upload a file."
    else:
        job = Job()
        job.source = {"type": "file" if up else "youtube", "value": chosen}
        ss.update(job=job, result=None, messages=[], warn=None)
        threading.Thread(target=_worker, args=(job, chosen, "english"), daemon=True).start()

def halt():
    job = st.session_state.get("job")
    if job:
        job.cancel.set()
        job.status = "stopped"

with st.container(border=True):
    tab_url, tab_file = st.tabs(["YouTube link", "Upload a file"])
    with tab_url:
        url = st.text_input("Video URL", placeholder="https://www.youtube.com/watch?v=…",
                            label_visibility="collapsed", disabled=running, key="url")
    with tab_file:
        f = st.file_uploader("Audio or video file", type=["mp3", "wav", "m4a", "mp4", "mkv", "webm", "mov"],
                             label_visibility="collapsed", disabled=running, key="up")
        if f:
            key = (f.name, f.size)
            if st.session_state.get("_up_key") != key:
                st.session_state["_up_key"], st.session_state["_up_path"] = key, save_upload(f)
    btn_col, _ = st.columns([1, 2])
    if running:
        btn_col.button("Stop Analysis", key="stop", on_click=halt)
    else:
        btn_col.button("Start Analysis", type="primary", key="start", on_click=begin)
    slot = st.empty()

if st.session_state.get("warn"):
    st.warning(st.session_state.warn)

job = st.session_state.get("job")
if job and job.status == "done" and not job.collected:
    job.result["title"] = clean_title(job.result["title"])
    job.result["source"] = job.source
    st.session_state.result, st.session_state.took, job.collected = job.result, job.took, True
    try:
        st.session_state.aid = save_record(job.result, job.took)
    except Exception:
        st.session_state.aid = "unsaved"
    st.toast("Analysis complete", icon="✅")
running = bool(job and job.status == "running")

if running:
    slot.markdown(progress_html(job.progress, job.stage, time.time() - job.started), unsafe_allow_html=True)
elif job and job.status == "stopped":
    slot.info("Analysis stopped. Start again whenever you like.")
elif job and job.status == "error":
    with slot.container():
        st.error(job.error)
        with st.expander("Technical details"):
            st.code(job.detail, language=None)

r = st.session_state.result
state = "run" if running else "done" if r else (job.status if job and job.status in ("stopped", "error") else "idle")
nav_slot.markdown(nav_html(state), unsafe_allow_html=True)

if not r and running:
    pass
elif not r:
    st.markdown(
        """
<div class="steps">
  <div class="step"><b>1</b><h4>Add your video</h4><p>Paste a YouTube link or upload a recording from your device.</p></div>
  <div class="step"><b>2</b><h4>Get the essentials</h4><p>A title, summary, action items, decisions and open questions.</p></div>
  <div class="step"><b>3</b><h4>Ask anything</h4><p>Chat with the transcript to find exactly what was said.</p></div>
</div>""",
        unsafe_allow_html=True,
    )
else:
    aid, src, msgs = st.session_state.aid, r.get("source") or {}, st.session_state.messages
    st.markdown("<div style='height:1.4rem'></div>", unsafe_allow_html=True)
    words = len(r["transcript"].split())
    mins = max(1, round(words / 150))
    took_m, took_s = divmod(int(st.session_state.took), 60)
    left, right = st.columns([3, 2], gap="large")
    with left:
        st.markdown(
            f'<div class="reveal"><p class="result-title" style="font-size:1.6rem; font-weight:700; margin:0 0 .9rem;">{html.escape(r["title"])}</p><div class="kpis">'
            f'<div class="kpi"><small>Speech</small><b>~{mins} min</b></div>'
            f'<div class="kpi"><small>Words</small><b>{words:,}</b></div>'
            f'<div class="kpi"><small>Analysed in</small><b>{took_m:02d}:{took_s:02d}</b></div></div></div>',
            unsafe_allow_html=True,
        )
        if src.get("value"):
            audio_only = src.get("type") == "file" and Path(src["value"]).suffix.lower() in (".mp3", ".wav", ".m4a")
            try:
                (st.audio if audio_only else st.video)(src["value"])
            except Exception:
                st.caption("Preview isn't available for this source.")
        with st.container(border=True):
            t1, t2, t3, t4, t5 = st.tabs(["Summary", "Action items", "Decisions", "Open questions", "Transcript"])
            
            for tab, key in ((t1, "summary"), (t3, "key_decisions"), (t4, "open_questions")):
                with tab:
                    render_copy_icon(r[key], key)
                    st.markdown(r[key])
            
            with t2:
                render_copy_icon(r["action_items"], "action_items")
                items = parse_items(r["action_items"])
                if items:
                    bar = st.empty()
                    done = sum(st.checkbox(it, key=f"ai_{aid}_{i}") for i, it in enumerate(items))
                    bar.progress(done / len(items), text=f"{done} of {len(items)} completed")
                else:
                    st.markdown(r["action_items"])
                    
            with t5:
                render_copy_icon(r["transcript"], "transcript")
                st.markdown(f'<div class="transcript-scroll-area">{html.escape(r["transcript"])}</div>', unsafe_allow_html=True)
                
        x = st.columns([1, 1, 1.4, 1.6])
        i = 0
        if HAS_DOCX:
            x[i].download_button("Word", build_docx(r["title"], sections_of(r), r["transcript"]), file_name="meeting_report.docx",
                                 mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
            i += 1
        if HAS_PDF:
            x[i].download_button("PDF", build_pdf(r["title"], sections_of(r), r["transcript"]),
                                 file_name="meeting_report.pdf", mime="application/pdf", use_container_width=True)
        if x[3].button("New Analysis", key="new", use_container_width=True):
            st.session_state.update(result=None, job=None, messages=[])
            st.rerun()
    with right:
        st.markdown('<div class="chat-head">Ask about this video</div>'
                    '<div class="chat-sub">Answers come from the transcript.</div>', unsafe_allow_html=True)
        box = st.container(height=470, border=True)
        with box:
            if not msgs:
                st.caption("Try: “Who owns the next steps?” or “What was decided about the deadline?”")
            for m in msgs:
                with st.chat_message(m["role"]):
                    st.markdown(m["content"])
                    show_sources(m.get("sources"))
        q = st.chat_input("Ask a question")
        if q:
            from core.rag_engine import ask_question, build_rag_chain
            msgs.append({"role": "user", "content": q})
            with box:
                with st.chat_message("user"):
                    st.markdown(q)
                with st.chat_message("assistant"):
                    try:
                        if r.get("rag_chain") is None:
                            with st.spinner("Preparing the chat index…"):
                                r["rag_chain"] = build_rag_chain(r["transcript"])
                        with st.spinner("Searching the transcript…"):
                            ans = ask_question(r["rag_chain"], q)
                    except Exception as e:
                        ans = f"Sorry, I couldn't answer that. {friendly_error(e)}"
                    try:
                        ans = st.write_stream(typewriter(ans))
                    except AttributeError:
                        st.markdown(ans)
                    srcs = find_sources(f"{q} {ans}", passages(r["transcript"]))
                    show_sources(srcs)
            msgs.append({"role": "assistant", "content": ans, "sources": srcs})

if running:
    time.sleep(0.4)
    st.rerun()